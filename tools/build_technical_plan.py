from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path('原理图检查技术方案.docx')
BLUE = '2E74B5'
DARK = '1F4D78'
NAVY = '0B2545'
LIGHT = 'E8EEF5'
GRAY = 'F2F4F7'
MUTED = '666666'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_cell_width(cell, width):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'), str(width)); tcW.set(qn('w:type'), 'dxa')

def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'), 'dxa')
    ind = tblPr.find(qn('w:tblInd'))
    if ind is None:
        ind = OxmlElement('w:tblInd'); tblPr.append(ind)
    ind.set(qn('w:w'), '120'); ind.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol'); col.set(qn('w:w'), str(width)); grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx]); set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_run_font(run, name='Microsoft YaHei', size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    for attr in ['ascii', 'hAnsi', 'eastAsia']:
        rFonts.set(qn('w:' + attr), name)
    if size is not None: run.font.size = Pt(size)
    if color is not None: run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def style_doc(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(1); sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1); sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492); sec.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles['Normal']; normal.font.name = 'Microsoft YaHei'; normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in [('Heading 1',16,BLUE,16,8),('Heading 2',13,BLUE,12,6),('Heading 3',11.5,DARK,8,4)]:
        st = styles[name]; st.font.name='Microsoft YaHei'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
    if 'Code Block' not in styles:
        st = styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
    else: st = styles['Code Block']
    st.font.name='Consolas'; st.font.size=Pt(9); st.font.color.rgb=RGBColor.from_string(NAVY)
    st.paragraph_format.left_indent=Inches(0.2); st.paragraph_format.right_indent=Inches(0.2); st.paragraph_format.space_before=Pt(3); st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.0
    # Running header/footer
    hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run('原理图自动检查技术方案'); set_run_font(r, size=8.5, color=MUTED)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run('DSNCheck | 技术方案'); set_run_font(r, size=8.5, color=MUTED)

def add_title(doc):
    p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(28); p.paragraph_format.space_after=Pt(6)
    r=p.add_run('原理图自动检查技术方案'); set_run_font(r,size=25,color=NAVY,bold=True)
    p = doc.add_paragraph(); p.paragraph_format.space_after=Pt(18)
    r=p.add_run('Cadence Capture 起步、面向多 EDA 扩展的离线数据检查架构'); set_run_font(r,size=13,color=MUTED)
    for label, value in [('文档定位','技术方案与实施参考'),('当前实现目录',r'D:\GPTPRJ\DSNCheck'),('覆盖范围','数据导出、数据包、Netlist、离线规则检查、多 EDA 适配'),('日期','2026-08-29')]:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
        a=p.add_run(label + '：'); set_run_font(a,size=10.5,bold=True)
        b=p.add_run(value); set_run_font(b,size=10.5)
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(12)
    pPr=p._p.get_or_add_pPr(); pb=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'10'); bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),BLUE); pb.append(bottom); pPr.append(pb)

def add_table(doc, headers, rows, widths=None):
    table=doc.add_table(rows=1, cols=len(headers)); table.style='Table Grid'
    hdr=table.rows[0].cells
    for i,h in enumerate(headers):
        hdr[i].text=''; p=hdr[i].paragraphs[0]; r=p.add_run(h); set_run_font(r,size=9.5,bold=True,color=NAVY); set_cell_shading(hdr[i],LIGHT)
    for row in rows:
        cells=table.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=''; p=cells[i].paragraphs[0]; r=p.add_run(str(val)); set_run_font(r,size=9)
    set_table_geometry(table, widths or [9360//len(headers)]*len(headers))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return table

def add_bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.08
        r=p.add_run(item); set_run_font(r,size=10.5)

def add_code(doc, text):
    p=doc.add_paragraph(style='Code Block'); p.add_run(text)

doc=Document(); style_doc(doc); add_title(doc)

doc.add_heading('1. 方案结论', level=1)
doc.add_paragraph('本方案采用“EDA 专用导出适配器 + 统一中间数据格式 + 离线检查引擎”的分层架构。Capture 只负责把当前工程导出为数据文件；检查器只读取数据文件，不再直接操作 DSN 或依赖 Capture 会话。该方式既能覆盖 Cadence 的原理图几何检查，也能为 KiCad、Altium、Eagle、Mentor/Siemens、PADS 等 EDA 保留扩展空间。')
add_bullets(doc, ['网络拓扑检查优先使用 ISCF、PSpice/逻辑 Netlist 或统一后的 nets/pins 数据。','悬空导线、Junction、Bus、Off-page Connector 和图形重叠必须保留 Wire、坐标、连接点及图形边界框。','检查规则只读取统一 JSON/JSONL/CSV 数据，不直接调用某个 EDA 的 API。','每个适配器必须报告来源 EDA、版本、导出能力和未支持字段，避免产生“未检查即通过”的误判。'])

doc.add_heading('2. 总体架构与运行边界', level=1)
add_code(doc, 'EDA 原理图 → 专用导出适配器 → 离线数据包 → 通用检查引擎 → TXT/JSON 报告')
add_table(doc, ['层次','职责','是否接触 DSN/EDA 会话'], [
    ('导出层','调用原生 API、Tcl/Dbo 或读取开放格式，生成原始数据','是，仅在导出阶段'),
    ('规范化层','将原始记录整理为统一的器件、Pin、Net、Wire、Sheet 等模型','否'),
    ('检查层','执行网络、属性、层次结构和几何规则','否'),
    ('报告层','按类别输出错误、位置、来源对象和证据','否'),
], [1500,5100,2760])
doc.add_paragraph('边界原则：导出阶段允许使用 Capture；离线检查阶段不读取 DSN、不连接 Capture、不修改原理图。导出的数据包可以归档、复查、在 CI 中重复检查。')

doc.add_heading('3. Cadence Capture 数据导出方案', level=1)
doc.add_heading('3.1 原生 ISCF', level=2)
doc.add_paragraph('Capture 16.6 自带 capISCFExport。批处理方式使用 ExportDesignInBatch；当前活动设计方式使用 ExportDesign $design，从而不重新打开 DSN、不创建独立 Dbo Session。ISCF 主要提供元件属性、RefDes、Pin、普通网络、Power、Ground 和 Bus 逻辑信息。')
add_code(doc, '::capISCFExport::ExportDesign $design $iscfPath $logPath')
doc.add_heading('3.2 Capture Tcl/Dbo 补充数据', level=2)
doc.add_paragraph('Dbo 导出器遍历设计、原理图、页面和页面对象，输出 JSONL。当前字段覆盖：')
add_bullets(doc, ['Part、Value、Reference、PCB Footprint、Source Library/Part；','Pin、Pin Number、Pin Name、Pin Type、No Connect、所属器件和网络；','Wire、Wire Point、Wire 端点对象数、网络名、方向、长度属性；','PointIsJunction、JunctionOnWire、Net Alias；','Global、Power、Ground、Port、Off-page Connector；','Bus Entry、Bus/Bundle 标识；','Graphic、Display Property、User Property、坐标和 Bounding Box。'])
doc.add_heading('3.3 当前 Capture 入口', level=2)
add_code(doc, '在 Capture Tcl Command Window 执行：\nsource {D:/GPTPRJ/DSNCheck/scripts/capture_export_active.tcl}')
doc.add_paragraph('默认输出到活动 DSN 同目录的 <DSN文件名>_checkdata。该入口只做导出；导出完成后可关闭 Capture，再执行后续离线步骤。')

doc.add_heading('4. 统一离线数据包', level=1)
add_table(doc, ['文件/目录','内容','主要用途'], [
    ('native/design.iscf','Cadence 原生 ISCF','属性、Pin、Net、Power/Ground、Bus'),
    ('dbo/objects.jsonl','完整 Dbo 对象记录','Wire、Junction、坐标、层次对象、图形'),
    ('dbo/properties.jsonl','User/Display Property','属性完整性和文字/图形位置'),
    ('dbo/errors.jsonl','导出阶段非致命错误','数据质量和兼容性追踪'),
    ('netlist/parts.csv','规范化器件表','器件、位号、值、封装'),
    ('netlist/pins.csv','规范化 Pin 表','Pin 与网络、器件的关系'),
    ('netlist/nets.csv','规范化网络汇总','连接数、器件数、电源标识'),
    ('netlist/logical_netlist.jsonl','逐网络结构化记录','通用网络规则输入'),
    ('package_manifest.json','数据包清单和能力信息','版本、统计、来源和完整性'),
], [2200,3600,3560])
doc.add_paragraph('Cadence PCB Netlist 的 pstxnet.dat、pstxprt.dat、pstchip.dat 可以作为可选外部输入收集；Cadence 16.6 的 Tools → Create Netlist → PCB Editor 没有发现稳定的公开 Tcl 调用接口。Cadence DRC 报告不纳入本方案的数据包。')

doc.add_heading('5. Netlist 的作用与限制', level=1)
doc.add_paragraph('Netlist 适合表达“最终逻辑连接结果”，不等价于完整原理图。它可以支持器件属性、位号/封装、Pin 连接、单节点网络、伪单节点网络和孤立电源；但通常无法证明导线在图面上是否真正连接，也无法表达 Junction、图形重叠或 Bus Entry 的几何关系。')
add_table(doc, ['检查项目','Netlist','完整 Dbo/几何数据'], [
    ('元件属性、位号、封装','支持','支持'),('引脚连接、单节点网络','支持','支持'),('孤立电源、伪单节点网络','支持（需器件类型/电源标识）','支持'),('悬空导线','不可靠','必须'),('Junction、导线交叉','不支持','必须'),('Off-page、Bus Entry','不完整','必须'),('图形/文字重叠','不支持','必须'),
], [2900,1800,4660])

doc.add_heading('6. 当前三类网络错误规则', level=1)
add_table(doc, ['类别','规则定义','所需字段','输出证据'], [
    ('单节点网络','有效连接 Pin 数少于 2 个','net_name、pin_id、owner_id、no_connect','网络名、连接数、Pin、器件'),
    ('孤立电源','电源网络的连接器件全部为电阻/电容','电源标识、RefDes、器件类型、Net','网络名、电源标识、R/C 器件'),
    ('伪单节点网络','网络至少有 2 个 Pin，但所有 Pin 属于同一器件','net_name、owner_id、Pin','网络名、器件、Pin 清单'),
], [1800,3300,2300,1960])
doc.add_paragraph('当前离线检查器读取 dbo/objects.jsonl，并输出 dsn_check_report.txt 和 dsn_check_report.json。规则中 No Connect Pin 不计入有效连接。电源识别优先使用 Global/Power 信息，同时兼容 GND、VCC、VDD、VSS、VBAT 等约定命名。')

doc.add_heading('7. 多 EDA 兼容策略', level=1)
add_table(doc, ['EDA','首选输入','适配器能力重点','几何数据可得性'], [
    ('Cadence Capture','Tcl/Dbo + ISCF','完整对象和网络','高'),('KiCad','.kicad_sch S-expression','直接解析开放格式','高'),('Altium','官方 API/脚本/导出 Netlist','器件、Pin、Net、属性；几何依赖 API','中-高'),('Eagle','XML .sch','器件、网络、图形','高'),('Mentor/Siemens','官方 API/ASCII/XML','版本适配和能力声明','依版本'),('PADS','ASCII/数据库/Netlist','器件和网络；几何依赖导出方式','中'),
], [1700,2500,3160,2000])
doc.add_paragraph('统一模型建议包含 design、sheet、component、pin、net、wire、wire_point、junction、net_alias、power_symbol、ground_symbol、port、offpage_connector、bus、bus_entry、graphic、property，并保留 source_eda、source_file、source_object_id、sheet_name、coordinates。')

doc.add_heading('8. 工具目录与使用流程', level=1)
add_table(doc, ['文件','用途'], [
    ('scripts/capture_export_active.tcl','从当前 Capture 活动设计导出数据'),
    ('scripts/capture_export_all.tcl','导出器核心库，含 ISCF/Dbo/自定义补充逻辑'),
    ('Export-DSNCheckPackage.ps1','将原始导出整理为离线数据包'),
    ('Run-DSNCheck.ps1','只读取离线数据包并生成检查报告'),
    ('Invoke-DSNCheck.ps1','网络规则检查核心实现'),
], [3000,6360])
doc.add_heading('8.1 标准两阶段流程', level=2)
add_code(doc, '1) Capture Tcl Window:\n   source {D:/GPTPRJ/DSNCheck/scripts/capture_export_active.tcl}\n\n2) PowerShell:\n   & D:\GPTPRJ\DSNCheck\Export-DSNCheckPackage.ps1 `\n     -InputDirectory D:\GPTPRJ\EXAMPLE_SCH\<DSN>_checkdata\n\n3) PowerShell:\n   & D:\GPTPRJ\DSNCheck\Run-DSNCheck.ps1 `\n     -InputDirectory D:\GPTPRJ\EXAMPLE_SCH\<DSN>_checkdata_package')
doc.add_paragraph('若指定 OutputDirectory，应使用纯 ASCII 路径；Capture 16.6 内置 Tcl 8.4 对中文路径兼容性较差。')

doc.add_heading('9. 当前验证结果与已知限制', level=1)
doc.add_paragraph('对示例工程 AT91SAM9M10-G45-EK_REVA2-TOOL-TEST 的一次导出和离线检查已完成：')
add_table(doc, ['指标','结果'], [('Dbo 对象数据','objects.jsonl 约 5.1 MB'),('器件','709'),('已连接 Pin','2396'),('网络','650'),('单节点网络','20'),('孤立电源','0'),('伪单节点网络','0')], [3000,6360])
doc.add_paragraph('报告中已列出 20 个单节点网络的网络名、连接数、器件和 Pin。当前检查器仍属于规则原型：孤立电源识别依赖 Power/Global 标识和命名约定；若某 EDA 未导出电源类型，应报告“无法判定”，而不是默认无错误。')
doc.add_heading('9.1 网表导出失败案例', level=2)
doc.add_paragraph('示例工程的 Capture PCB Netlist 日志显示 ORCAP-36071：J1-1 至 J1-4 的 PCB Footprint 属性包含空格，随后 ORCAP-36018 中止网表生成。SH、SUP、LCD、M、Z 等对象的“无引脚”信息是警告，不是中止根因。')

doc.add_heading('10. 后续实施路线', level=1)
add_bullets(doc, ['固化统一 schema_version、字段类型、坐标单位和错误码。','为每个 EDA 适配器增加能力清单与最小样例工程回归测试。','扩展几何规则：悬空导线、Junction、导线相交、Off-page 匹配、Bus Entry、图形重叠。','让检查器同时读取 ISCF、规范化 Netlist 和 Dbo JSONL，并对字段冲突给出证据。','增加 HTML/JSON 报告、按页定位、源对象 ID 和可选图片标注。','把离线数据包接入批处理/CI，固定输入快照，支持规则版本和结果对比。'])

doc.add_heading('附录 A：设计决策摘要', level=1)
add_table(doc, ['决策','结论'], [
    ('是否直接解析 Cadence DSN','不解析二进制 DSN，使用官方 ISCF 与 Tcl/Dbo'),
    ('是否让检查器连接 Capture','不连接；检查器只读离线数据包'),
    ('Netlist 是否足够','足够做网络/器件逻辑检查，不足以做几何检查'),
    ('是否纳入 Cadence DRC','本方案数据包不纳入 DRC 报告'),
    ('多 EDA 如何扩展','每种 EDA 一个适配器，输出统一中间模型'),
], [3000,6360])
doc.add_paragraph('文档结束。')

doc.save(OUT)
print(OUT.resolve())
